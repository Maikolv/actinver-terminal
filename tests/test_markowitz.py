"""Constructor de portafolios (Markowitz), Excel con fórmulas y clasificación por capitalización bursátil."""
import io

import numpy as np
import pandas as pd
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from terminal import capitalizacion, db, excel_markowitz, markowitz, ranking, vigencia

from conftest import sembrar_precios

IDS = ["SIC:AAPL", "SIC:MU", "SIC:NVDA", "BMV:AMX", "BMV:WALMEX"]


def sembrar(con):
    sembrar_precios(con, IDS, fin=vigencia.ultima_sesion_cerrada("XNYS"), sesiones=320, proveedor="archivo")


def test_categorias_y_lectura_sec():
    assert [capitalizacion.categoria(x) for x in (3e12, 50e9, 5e9, 1e9, 1e8, None)] == ["mega", "large", "mid", "small", "micro", None]
    datos = {"units": {"shares": [
        {"val": 100, "end": "2025-01-01", "filed": "2025-02-01", "accn": "a", "form": "10-K"},
        {"val": 60, "end": "2026-06-30", "filed": "2026-08-01", "accn": "b", "form": "10-Q", "frame": "CY2026Q2I"},
        {"val": 60, "end": "2026-06-30", "filed": "2026-08-01", "accn": "b", "form": "10-Q"},  # misma cifra repetida
        {"val": 40, "end": "2026-06-30", "filed": "2026-08-01", "accn": "b", "form": "10-Q"}]}}
    assert capitalizacion.acciones_sec(datos) == (100, "2026-06-30", "10-Q")
    assert capitalizacion.flotante_sec({"units": {"USD": [{"val": 5e9, "end": "2025-06-30", "filed": "2025-08-01"}]}}) == (5e9, "2025-06-30")


def test_capitalizacion_usa_el_flotante_si_la_sec_solo_reporta_una_serie(con, ajustes):
    sembrar(con)
    with db.transaccion(con):
        capitalizacion._guardar(con, "SIC:AAPL", 15e9, "2026-06-30", "10-Q", "sec_edgar", "verificado", "ok")
        capitalizacion._guardar(con, "SIC:MU", 1e6, "2026-06-30", "10-Q", "sec_edgar", "verificado", "una serie", 500e9, "2025-06-30")
        capitalizacion._guardar(con, "SIC:NVDA", 9e9, "2026-06-30", "20-F", "sec_edgar", "extranjero", "ADR")
    fx = 17.0
    caps = capitalizacion.para_precios(con, ajustes, {"SIC:AAPL": 200 * fx, "SIC:MU": 100 * fx, "SIC:NVDA": 10 * fx})
    assert caps["SIC:AAPL"]["cap_categoria"] == "mega" and caps["SIC:AAPL"]["cap_estado"] == "verificado"
    assert caps["SIC:MU"]["cap_estado"] == "aproximado" and caps["SIC:MU"]["cap_usd"] == 500e9
    assert caps["SIC:NVDA"]["cap_categoria"] is None  # emisor extranjero: no se clasifica


def test_markowitz_metodos_respetan_restricciones_y_teoria(con, ajustes):
    sembrar(con)
    base = {"activos": IDS, "peso_max": 0.4, "rf": 0.05, "ventana": 252}
    res = {m: markowitz.calcular(con, ajustes, {**base, "metodo": m}) for m in markowitz.METODOS}
    for r in res.values():
        w = np.array([a["peso"] for a in r["activos"]])
        assert abs(w.sum() - 1) < 1e-6 and w.min() >= -1e-9 and w.max() <= 0.4 + 1e-6
    t, mv, eq = res["max_sharpe"]["metricas"], res["min_varianza"]["metricas"], res["equiponderado"]["metricas"]
    assert t["sharpe"] >= eq["sharpe"] - 1e-6 and t["sharpe"] >= mv["sharpe"] - 1e-6
    assert mv["riesgo"] <= min(eq["riesgo"], t["riesgo"]) + 1e-6
    fr = res["max_sharpe"]["frontera"]
    assert len(fr) >= 10 and all(b["rend"] >= a["rend"] - 1e-9 for a, b in zip(fr, fr[1:]))
    objetivo = (fr[0]["rend"] + fr[-1]["rend"]) / 2
    r = markowitz.calcular(con, ajustes, {**base, "metodo": "rend_objetivo", "rend_objetivo": objetivo})
    assert abs(r["metricas"]["rend"] - objetivo) < 1e-4
    assert sum(a["titulos"] * a["precio_mxn"] for a in r["activos"]) <= r["especificacion"]["capital"] + 1e-6


@pytest.mark.parametrize("spec,mensaje", [
    ({"activos": ["SIC:AAPL"]}, "de 2 a"),
    ({"activos": ["SIC:AAPL", "NO:EXISTE"]}, "desconocidos"),
    ({"activos": IDS, "peso_max": 0.1}, "no se llega a 100"),
    ({"activos": IDS, "metodo": "magia"}, "método"),
])
def test_markowitz_valida_la_especificacion(con, ajustes, spec, mensaje):
    sembrar(con)
    with pytest.raises(markowitz.ErrorPortafolio, match=mensaje):
        markowitz.calcular(con, ajustes, spec)


def test_persistencia_y_excel_con_formulas(con, ajustes):
    sembrar(con)
    p = markowitz.guardar(con, ajustes, {"nombre": "Retiro", "activos": IDS, "objetivo": "autorrealizacion"})
    markowitz.guardar(con, ajustes, {"nombre": "Retiro 2", "activos": IDS[:3], "metodo": "min_varianza", "peso_max": 0.5}, p["id"])
    assert [x["nombre"] for x in markowitz.listar(con)] == ["Retiro 2"]
    wb = load_workbook(io.BytesIO(excel_markowitz.libro(con, ajustes, markowitz.listar(con))))
    assert wb.sheetnames[:4] == ["Guía", "Resumen", "P1 Datos", "P1 Modelo"]
    modelo = wb["P1 Modelo"]
    formulas = [c.value for fila in modelo.iter_rows() for c in fila if isinstance(c.value, str) and c.value.startswith("=")]
    assert any("COVARIANCE.S" in f for f in formulas) and any(f.startswith("=SQRT(MAX(") for f in formulas)
    assert wb["Resumen"]["G4"].value.startswith("='P1 Modelo'!")
    plantilla = load_workbook(io.BytesIO(excel_markowitz.plantilla()))
    assert plantilla.sheetnames == ["Guía", "Paramétrico Acelera", "Histórico Datos", "Histórico Modelo"]
    assert plantilla["Guía"]._images and plantilla["Histórico Modelo"]._images  # logo de la terminal en cada hoja
    assert not any(str(c.value).startswith("=") and "#REF!" in c.value for ws in plantilla for fila in ws.iter_rows() for c in fila)
    assert markowitz.eliminar(con, p["id"]) and not markowitz.listar(con)


def test_api_portafolios_exige_csrf_y_descarga_excel(monkeypatch):
    from terminal.app import app
    from terminal.seguridad import TOKEN_CSRF
    c0 = db.conectar()
    db.inicializar(c0)
    sembrar(c0)
    c0.close()
    h = {"X-CSRF-Token": TOKEN_CSRF}
    spec = {"nombre": "API", "activos": IDS[:4]}
    with TestClient(app, base_url="http://127.0.0.1:8765") as c:
        assert c.post("/api/portafolios", json=spec).status_code == 403
        assert c.post("/api/portafolios/calcular", json={"activos": ["SIC:AAPL"]}, headers=h).status_code == 422
        assert c.post("/api/portafolios/generar", json={"completar": False, "fijos": IDS[:3]}).status_code == 403
        g = c.post("/api/portafolios/generar", json={"completar": False, "fijos": IDS[:3]}, headers=h)
        assert g.status_code == 200 and len(g.json()["resultado"]["activos"]) == 3
        assert c.post("/api/portafolios/generar", json={"objetivo": "otro"}, headers=h).status_code == 422
        pid = c.post("/api/portafolios", json=spec, headers=h).json()["id"]
        assert any(p["id"] == pid for p in c.get("/api/portafolios").json()["portafolios"])
        x = c.get(f"/api/portafolios/{pid}/excel")
        assert x.status_code == 200 and x.headers["content-type"].startswith("application/vnd.openxml")
        assert c.get("/api/markowitz/plantilla.xlsx").status_code == 200
        assert c.request("DELETE", f"/api/portafolios/{pid}", json={}, headers=h).status_code == 200
        assert c.get(f"/api/portafolios/{pid}").status_code == 404


def test_generador_completa_con_el_ranking_y_sugiere(con, ajustes):
    from terminal import mercado
    sembrar(con)
    ranking._CACHE.clear()
    r = markowitz.generar(con, ajustes, {"objetivo": "crecimiento", "n": 4, "fijos": ["SIC:MU"], "excluir": ["SIC:AAPL"]})
    ids = [a["id"] for a in r["resultado"]["activos"]]
    assert ids[0] == "SIC:MU" and "SIC:AAPL" not in ids and len(ids) == 4
    assert r["fijos"] == ["SIC:MU"] and set(r["sugeridos"]) == set(ids[1:]) and all(r["motivos"][i] for i in r["sugeridos"])
    assert all(s["texto"] and (not s.get("accion") or s.get("etiqueta")) for s in r["sugerencias"])
    # sin completar usa solo lo fijado; un tope que no llega a 100 % se sube y se avisa
    r2 = markowitz.generar(con, ajustes, {"objetivo": "conservador", "completar": False, "fijos": IDS[:3], "peso_max": 0.1})
    assert [a["id"] for a in r2["resultado"]["activos"]] == IDS[:3] and r2["sugeridos"] == []
    assert r2["resultado"]["especificacion"]["peso_max"] == 0.34 and r2["sugerencias"][0]["tipo"] == "aviso"
    for malo in ({"objetivo": "maslow"}, {"fijos": ["NO:EXISTE"]}, {"clase": "bonos"}, {"caps": ["gigante"]}):
        with pytest.raises(markowitz.ErrorPortafolio):
            markowitz.generar(con, ajustes, malo)
    # los portafolios guardados con los nombres anteriores de objetivo se siguen leyendo
    assert markowitz.normalizar({"objetivo": "seguridad", "activos": IDS[:2]}, mercado.instrumentos(con))["objetivo"] == "conservador"


def test_ranking_en_cache_hasta_que_llegan_datos(con, ajustes, monkeypatch):
    sembrar(con)
    ranking._CACHE.clear()
    llamadas = []
    original = ranking.calcular
    monkeypatch.setattr(ranking, "calcular", lambda *a: llamadas.append(1) or original(*a))
    monkeypatch.setattr(ranking, "TOLERANCIA_S", 0)
    a = ranking.calcular_cache(con, ajustes)
    b = ranking.calcular_cache(con, ajustes)
    assert len(llamadas) == 1 and a["cache_s"] == 0 and b["filas"] == a["filas"]  # datos iguales: cálculo anterior
    with db.transaccion(con):  # un precio nuevo cambia la marca y obliga a recalcular
        con.execute("INSERT OR REPLACE INTO precios SELECT * FROM precios ORDER BY rowid DESC LIMIT 1")
    ranking.calcular_cache(con, ajustes)
    assert len(llamadas) == 2


def test_tasa_libre_de_cetes_y_respaldo(con, ajustes, monkeypatch):
    import httpx
    monkeypatch.setattr(markowitz, "credencial", lambda p: "token-de-prueba")
    ok = httpx.Client(transport=httpx.MockTransport(lambda r: httpx.Response(
        200, json={"bmx": {"series": [{"idSerie": "SF43936", "datos": [{"fecha": "29/10/2026", "dato": "6.22"}]}]}})))
    falla = httpx.Client(transport=httpx.MockTransport(lambda r: httpx.Response(503)))
    t = markowitz.tasa_libre(con, ajustes, falla)
    assert t["valor"] == 0.07 and t["fecha"] is None and "7 % fijo" in t["aviso"]          # nunca hubo dato: lo dice
    t = markowitz.tasa_libre(con, ajustes, ok)
    assert t["valor"] == 0.0622 and t["fecha"] == "2026-10-29" and "SF43936" in t["fuente"] and not t["aviso"]
    assert markowitz.tasa_libre(con, ajustes, falla)["valor"] == 0.0622                    # dentro de 12 h: sin consultar
    con.execute("UPDATE ajustes_usuario SET valor=json_set(valor, '$.consultado_en', '2026-01-01T00:00:00+00:00') "
                "WHERE clave='cetes28'")
    t = markowitz.tasa_libre(con, ajustes, falla)
    assert t["valor"] == 0.0622 and "última tasa guardada" in t["aviso"]


def test_fuera_de_muestra_ledoit_wolf_y_topes_por_moneda(con, ajustes):
    sembrar(con)
    base = {"activos": IDS, "metodo": "max_sharpe", "peso_max": 0.6, "ventana": 126}
    r = markowitz.calcular(con, ajustes, base)
    f = r["fuera_de_muestra"]
    assert f["disponible"] and f["meses"] >= 6 and set(f) >= {"propuesta", "1/N", "lectura", "nota"}
    assert f["entrenamiento"] == 126 and f["propuesta"]["rotacion_mensual"] >= 0 and "1/N" in f["lectura"]
    lw = markowitz.calcular(con, ajustes, {**base, "covarianza": "ledoit_wolf"})
    assert lw["covarianza_nombre"].startswith("Ledoit") and lw["metricas"]["riesgo"] != r["metricas"]["riesgo"]
    usd = markowitz.calcular(con, ajustes, {**base, "usd_max": 0.3})
    peso_usd = sum(a["peso"] for a in usd["activos"] if a["id"].startswith("SIC:"))
    assert peso_usd <= 0.3 + 1e-4 and usd["restricciones"][0]["grupo"] == "USD (SIC)"
    with pytest.raises(markowitz.ErrorPortafolio, match="no se llega a 100"):   # 2 BMV × 30 % < 100 % sin USD
        markowitz.calcular(con, ajustes, {**base, "peso_max": 0.3, "usd_max": 0})
    with pytest.raises(markowitz.ErrorPortafolio):
        markowitz.calcular(con, ajustes, {**base, "covarianza": "otra"})
    corto = markowitz.fuera_de_muestra(pd.DataFrame({"a": [1.0] * 100, "b": [1.0] * 100}),
                                       {"ventana": 252, "covarianza": "muestral", "metodo": "equiponderado", "rf": 0.07},
                                       0, 1, [], None)
    assert not corto["disponible"] and "sesiones" in corto["motivo"]


def test_black_litterman_vistas_confianza_y_fuera_de_muestra(con, ajustes):
    sembrar(con)
    base = {"activos": IDS, "metodo": "max_sharpe", "peso_max": 0.6, "ventana": 126}
    baja = markowitz.calcular(con, ajustes, {**base, "vistas": [
        {"activo": "SIC:AAPL", "rend": 0.4, "confianza": 0.1}]})
    alta = markowitz.calcular(con, ajustes, {**base, "vistas": [
        {"activo": "SIC:AAPL", "rend": 0.4, "confianza": 0.9}]})
    assert alta["activos"][0]["rend"] > baja["activos"][0]["rend"]
    assert alta["fuera_de_muestra"]["disponible"] and "vistas de Black-Litterman son las de hoy" in alta["fuera_de_muestra"]["nota"]
    assert alta["fuera_de_muestra"]["propuesta"]["rend_total"] != baja["fuera_de_muestra"]["propuesta"]["rend_total"]
    generado = markowitz.generar(con, ajustes, {"completar": False, "fijos": IDS[:3], "vistas": [
        {"activo": "SIC:AAPL", "contra": "SIC:MU", "rend": 0.05, "confianza": 0.7}]})
    assert generado["resultado"]["especificacion"]["vistas"][0]["contra"] == "SIC:MU"
    with pytest.raises(markowitz.ErrorPortafolio, match="confianza"):
        markowitz.calcular(con, ajustes, {**base, "vistas": [{"activo": "SIC:AAPL", "rend": 0.1, "confianza": 0}]})


def test_sectores_por_fuente_sin_mezclar_taxonomias():
    import csv
    filas = list(csv.DictReader(markowitz.ARCHIVO_SECTORES.open(encoding="utf-8")))
    assert all(f["fuente"].strip() for f in filas)                           # cada fila dice de dónde sale o por qué falta
    s = markowitz.sectores()
    assert s["SIC:XLK"] == s["SIC:VGT"] == "ETF · Tecnología"                 # sector declarado en el nombre del fondo
    assert s["SIC:NVDA"].startswith("SIC ")                                  # acciones de EE. UU.: grupo SIC de la SEC
    assert "SIC:SPY" not in s and "BMV:WALMEX" not in s                      # amplios y BMV: sin sector, no se inventa
