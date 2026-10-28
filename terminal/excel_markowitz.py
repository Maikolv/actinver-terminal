"""Libros de Excel con el desarrollo del modelo de Markowitz, con fórmulas vivas (se pueden editar como plantilla).

* `plantilla()`: ejercicio paramétrico del Taller de Acelera (IPC, S&P 500 y Treasury MXN) con las fórmulas que en
  el original quedaron como #REF!, más una hoja histórica vacía para pegar precios propios.
* `libro(...)`: los portafolios guardados del generador, cada uno con sus precios, rendimientos, estadísticos,
  covarianzas, correlaciones, pesos, resultados, frontera eficiente y gráfica.

Llevan el diseño de la terminal (logo, colores y tipografía de `web/static/app.css`) y una hoja «Guía» con los pasos
y las celdas exactas de cada libro. La varianza del portafolio se calcula sin fórmulas matriciales:
(Σw)_i = Σ_j cov_ij·w_j y σ²_p = Σ_i w_i·(Σw)_i, así funciona igual en Excel, LibreOffice y Google Sheets.
"""
from __future__ import annotations

import io
from datetime import date
from functools import lru_cache

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.image import Image as ImagenXL
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

from . import markowitz

# Colores de la terminal (modo claro de app.css) para que el libro se vea como la terminal.
MARCA, TEXTO, TEXTO_2, BORDE = "1F5F8B", "1C1D1F", "55585E", "D9D6CF"
SUPERFICIE_2, POSITIVO, POSITIVO_FONDO, SERIE_2 = "EFEDE8", "1D6B3A", "E3F1E7", "B0600F"
FUENTE = "Segoe UI"
NEGRITA = Font(name=FUENTE, bold=True, color=TEXTO)
TITULO = Font(name=FUENTE, bold=True, size=15, color=TEXTO)
SECCION = Font(name=FUENTE, bold=True, size=12, color=MARCA)
SUAVE = Font(name=FUENTE, size=10, color=TEXTO_2)
CABECERA = PatternFill("solid", fgColor=SUPERFICIE_2)
ENTRADA = PatternFill("solid", fgColor="E3EEF6")      # celdas que el usuario puede cambiar (azul claro)
RESULTADO = PatternFill("solid", fgColor=POSITIVO_FONDO)
LINEA = Side(style="thin", color=BORDE)
PCT, NUM = "0.00%", "#,##0.00"
# Ejercicio de Acelera: E(r) y σ anuales; correlaciones derivadas de las covarianzas diarias del mismo archivo
ACELERA = (["IPC", "SP500", "Treasury MXN"], [0.11, 0.16, 0.09], [0.18, 0.17, 0.10],
           [[1, 0.324, -0.265], [0.324, 1, 0.472], [-0.265, 0.472, 1]])
AVISO = "Simulación educativa con datos históricos: no es pronóstico ni recomendación. La terminal no envía órdenes."


@lru_cache(maxsize=1)
def _logo_png(px: int = 96) -> bytes:
    """El mismo logo que web/static/favicon.svg (cuadro azul con línea de precio blanca), en PNG."""
    from PIL import Image, ImageDraw
    k = px / 32
    im = Image.new("RGBA", (px, px), (0, 0, 0, 0))
    d = ImageDraw.Draw(im)
    d.rounded_rectangle([0, 0, px - 1, px - 1], radius=round(6 * k), fill="#" + MARCA)
    w = round(2.6 * k)
    for trazo in ([(6, 25), (26, 25)], [(8, 20), (13, 14), (17, 17), (24, 8)]):
        puntos = [(x * k, y * k) for x, y in trazo]
        d.line(puntos, fill="white", width=w, joint="curve")
        for x, y in puntos:  # extremos y vértices redondeados, como stroke-linecap/linejoin «round»
            d.ellipse([x - w / 2, y - w / 2, x + w / 2, y + w / 2], fill="white")
    buf = io.BytesIO()
    im.save(buf, "PNG")
    return buf.getvalue()


def _encabezado(ws, titulo: str, columnas: int, subtitulo: str | None = None) -> None:
    """Fila 1 como la cabecera de la terminal: logo, nombre y título, con línea inferior."""
    logo = ImagenXL(io.BytesIO(_logo_png()))
    logo.width = logo.height = 30
    ws.add_image(logo, "A1")
    ws.row_dimensions[1].height = 32
    ws["A1"] = f"Actinver Terminal · {titulo}"
    ws["A1"].font = TITULO
    ws["A1"].alignment = Alignment(vertical="center", indent=3)
    for c in range(1, columnas + 1):
        ws.cell(1, c).border = Border(bottom=Side(style="medium", color=MARCA))
    if subtitulo:
        ws["A2"] = subtitulo
        ws["A2"].font = SUAVE
    ws.sheet_properties.tabColor = MARCA


def _cabecera(celda) -> None:
    celda.font, celda.fill, celda.border = NEGRITA, CABECERA, Border(bottom=LINEA)


def _entrada(celda, formato: str | None = None) -> None:
    celda.fill, celda.border = ENTRADA, Border(left=LINEA, right=LINEA, top=LINEA, bottom=LINEA)
    if formato:
        celda.number_format = formato


def _fuente_terminal(wb: Workbook) -> None:
    """Segoe UI en todo el libro (openpyxl deja Calibri en las celdas sin formato)."""
    for ws in wb.worksheets:
        for fila in ws.iter_rows():
            for c in fila:
                if c.font.name != FUENTE:
                    f = c.font
                    c.font = Font(name=FUENTE, bold=f.bold, italic=f.italic, size=f.size, color=f.color)


def _guia(wb: Workbook, titulo: str, bloques: list[tuple[str, list[str]]]) -> None:
    """Hoja «Guía»: pasos numerados, cortos y con las celdas exactas del libro."""
    ws = wb.active
    ws.title = "Guía"
    ws.sheet_view.showGridLines = False
    _encabezado(ws, titulo, 2, f"Generado el {date.today():%d/%m/%Y}. {AVISO}")
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 118
    ws["A2"].alignment = Alignment(indent=0)
    fila = 4
    for encabezado, pasos in bloques:
        ws.cell(fila, 1, encabezado).font = SECCION
        fila += 1
        for k, paso in enumerate(pasos, start=1):
            ws.cell(fila, 1, k if len(pasos) > 1 else "•").font = Font(name=FUENTE, bold=True, color=MARCA)
            ws.cell(fila, 1).alignment = Alignment(horizontal="right", vertical="top")
            c = ws.cell(fila, 2, paso)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[fila].height = 15 * max(1, -(-len(paso) // 115))
            fila += 1
        fila += 1
    ejemplo = [("Azul claro", ENTRADA, "usted escribe aquí"), ("Blanco", None, "fórmula: no la borre"),
               ("Verde", RESULTADO, "resultado principal")]
    ws.cell(fila, 1, "Colores del libro").font = SECCION
    for k, (nombre, relleno, uso) in enumerate(ejemplo, start=fila + 1):
        c = ws.cell(k, 2, f"{nombre}: {uso}")
        if relleno:
            c.fill = relleno
    ws.freeze_panes = "A4"


def _matriz(ws, fila: int, col: int, nombres: list[str], celda) -> tuple[int, int]:
    """Escribe una matriz n×n; celda(i, j) devuelve la fórmula. Regresa (fila, col) de la esquina superior izquierda."""
    _cabecera(ws.cell(fila, col))
    for k, nm in enumerate(nombres):
        _cabecera(ws.cell(fila, col + 1 + k, nm))
        _cabecera(ws.cell(fila + 1 + k, col, nm))
    for i in range(len(nombres)):
        for j in range(len(nombres)):
            ws.cell(fila + 1 + i, col + 1 + j, celda(i, j)).number_format = "0.000000"
    return fila + 1, col + 1


def _resultados(ws, fila: int, n: int, f_e: int, c_e: str, c_w: str, f_cov: int, c_cov: int, rf: str, a: str) -> dict:
    """Bloque (Σw)_i, contribución al riesgo y resultados del portafolio. Devuelve las celdas clave."""
    fin = f_e + n - 1
    cw = f"{c_w}{f_e}:{c_w}{fin}"
    fila_w = fila + n + 9  # fila auxiliar con los pesos en horizontal (para multiplicar por renglones de la matriz)
    ws.cell(fila_w, 1, "Pesos en fila (auxiliar)").font = SUAVE
    for j in range(n):
        ws.cell(fila_w, c_cov + j, f"={c_w}{f_e + j}").number_format = PCT
    fila_aux = f"{L(c_cov)}{fila_w}:{L(c_cov + n - 1)}{fila_w}"
    ws.cell(fila - 1, 1, "Resultados del portafolio").font = SECCION
    for k, t in enumerate(["Activo", "(Σw)_i", "Contribución al riesgo"]):
        _cabecera(ws.cell(fila, 1 + k, t))
    for i in range(n):
        r = fila + 1 + i
        ws.cell(r, 1, f"=A{f_e + i}")
        ws.cell(r, 2, f"=SUMPRODUCT({L(c_cov)}{f_cov + i}:{L(c_cov + n - 1)}{f_cov + i},{fila_aux})").number_format = "0.000000"
        ws.cell(r, 3, f"=IFERROR({c_w}{f_e + i}*B{r}/$B${fila + n + 3},\"\")").number_format = PCT
    r0 = fila + n + 2
    celdas = {}
    principales = {"E(r_p) anual", "σ_p anual", "Sharpe"}
    for k, (et, f, fmt) in enumerate([
        ("E(r_p) anual", f"=SUMPRODUCT({cw},{c_e}{f_e}:{c_e}{fin})", PCT),
        ("Varianza σ²_p", f"=SUMPRODUCT({cw},B{fila + 1}:B{fila + n})", "0.000000"),
        ("σ_p anual", f"=SQRT(MAX(B{r0 + 1},0))", PCT),
        ("Sharpe", f"=IFERROR((B{r0}-{rf})/B{r0 + 2},\"\")", "0.000"),
        ("Utilidad U = E − ½·A·σ²", f"=B{r0}-0.5*{a}*B{r0 + 1}", "0.0000"),
        ("Asignación al riesgoso y* = (E−r_f)/(A·σ²)", f"=IFERROR((B{r0}-{rf})/({a}*B{r0 + 1}),\"\")", PCT),
        ("Suma de pesos (debe ser 100 %)", f"=SUM({cw})", PCT),
    ]):
        ws.cell(r0 + k, 1, et).font = NEGRITA
        c = ws.cell(r0 + k, 2, f)
        c.number_format = fmt
        if et in principales:
            c.fill, c.font = RESULTADO, Font(name=FUENTE, bold=True, color=POSITIVO)
        celdas[et] = f"B{r0 + k}"
    celdas["pesos"] = cw
    return celdas


def _colorear(serie: Series, color: str, marcador: str | None = None, linea: bool = True) -> None:
    if linea:
        serie.graphicalProperties.line.solidFill = color
        serie.graphicalProperties.line.width = 28000
    else:
        serie.graphicalProperties.line.noFill = True
    if marcador:
        serie.marker.symbol, serie.marker.size = marcador, 7
        relleno = GraphicalProperties(solidFill=color)
        relleno.line.solidFill = color
        serie.marker.graphicalProperties = relleno


def _frontera(ws, fila: int, res: dict, ancla: str) -> None:
    ws.cell(fila, 1, "Frontera eficiente (calculada por la terminal; Solver la reproduce)").font = SECCION
    for col, t in ((1, "σ_p"), (2, "E(r_p)"), (4, "Activo"), (5, "σ"), (6, "E(r)"), (8, "Elegido σ"), (9, "Elegido E(r)")):
        _cabecera(ws.cell(fila + 1, col, t))
    for k, p in enumerate(res["frontera"]):
        ws.cell(fila + 2 + k, 1, p["riesgo"]).number_format = PCT
        ws.cell(fila + 2 + k, 2, p["rend"]).number_format = PCT
    for k, a in enumerate(res["activos"]):
        ws.cell(fila + 2 + k, 4, a["clave"])
        ws.cell(fila + 2 + k, 5, a["riesgo"]).number_format = PCT
        ws.cell(fila + 2 + k, 6, a["rend"]).number_format = PCT
    m = res["metricas"]
    ws.cell(fila + 2, 8, m["riesgo"]).number_format = PCT
    ws.cell(fila + 2, 9, m["rend"]).number_format = PCT
    nf, na = len(res["frontera"]), len(res["activos"])
    if not nf:
        return
    g = ScatterChart()
    g.title, g.style, g.height, g.width = "Frontera eficiente", 13, 9, 18
    g.x_axis.title, g.y_axis.title = "Riesgo σ (anual)", "Rendimiento esperado (anual)"
    for xs, ys, nombre, color, marcador, linea in [
        (Reference(ws, min_col=1, min_row=fila + 2, max_row=fila + 1 + nf), Reference(ws, min_col=2, min_row=fila + 2, max_row=fila + 1 + nf), "Frontera", MARCA, None, True),
        (Reference(ws, min_col=5, min_row=fila + 2, max_row=fila + 1 + na), Reference(ws, min_col=6, min_row=fila + 2, max_row=fila + 1 + na), "Activos", TEXTO_2, "circle", False),
        (Reference(ws, min_col=8, min_row=fila + 2), Reference(ws, min_col=9, min_row=fila + 2), "Elegido", SERIE_2, "diamond", False)]:
        se = Series(ys, xs, title=nombre)
        _colorear(se, color, marcador, linea)
        g.series.append(se)
    ws.add_chart(g, ancla)


def _hoja_historica(wb: Workbook, prefijo: str, nombres: list[str], precios: pd.DataFrame | None, filas_vacias: int,
                    pesos: list[float], rf: float, aversion: float, peso_max: float, peso_min: float, res: dict | None) -> dict:
    n = len(nombres)
    datos = wb.create_sheet(f"{prefijo}Datos"[:31])
    cr = n + 3  # primera columna de rendimientos
    _encabezado(datos, "Precios (MXN) y rendimientos diarios", cr + n - 1)
    _cabecera(datos.cell(2, 1, "Fecha"))
    for j, nm in enumerate(nombres):
        c = datos.cell(2, 2 + j, nm)
        _cabecera(c)
        if precios is None:
            _entrada(c)
        _cabecera(datos.cell(2, cr + j, f"r {nm}" if precios is not None else f"=\"r \"&{L(2 + j)}2"))
    total = len(precios) if precios is not None else filas_vacias
    for k in range(total):
        r = 3 + k
        if precios is not None:
            datos.cell(r, 1, precios.index[k].date()).number_format = "yyyy-mm-dd"
            for j, i in enumerate(precios.columns):
                datos.cell(r, 2 + j, float(precios.iloc[k, j])).number_format = NUM
        else:
            for j in range(n + 1):
                _entrada(datos.cell(r, 1 + j))
        if k:
            for j in range(n):
                p = L(2 + j)
                datos.cell(r, cr + j, f'=IFERROR({p}{r}/{p}{r - 1}-1,"")').number_format = "0.0000%"
    ultima = 2 + total
    datos.column_dimensions["A"].width = 13
    for c in range(2, cr + n):
        datos.column_dimensions[L(c)].width = 13
    datos.freeze_panes = "B3"
    rango = lambda j: f"'{datos.title}'!${L(cr + j)}$4:${L(cr + j)}${ultima}"  # noqa: E731  (fila 3 no tiene rendimiento)

    ws = wb.create_sheet(f"{prefijo}Modelo"[:31])
    _encabezado(ws, f"Modelo de Markowitz · {res['especificacion']['nombre'] if res else 'plantilla con sus datos'}", n + 9,
                "Celdas azul claro: usted escribe. Verde: resultado principal. " + AVISO)
    for k, (et, v, fmt) in enumerate([("Tasa libre de riesgo r_f", rf, PCT), ("Aversión al riesgo A", aversion, "0.0"),
                                      ("Peso máximo por activo", peso_max, PCT), ("Peso mínimo por activo", peso_min, PCT)]):
        ws.cell(3 + k, 1, et).font = NEGRITA
        _entrada(ws.cell(3 + k, 2, v), fmt)
    f_e = 10
    ws.cell(f_e - 2, 1, "Activos: estadísticos anuales y pesos").font = SECCION
    for k, t in enumerate(["Activo", "E(r) anual", "σ anual", "Sharpe", "Peso w", "Máx.", "Mín."]):
        _cabecera(ws.cell(f_e - 1, 1 + k, t))
    for i, nm in enumerate(nombres):
        r = f_e + i
        ws.cell(r, 1, nm if precios is not None else f"='{datos.title}'!{L(2 + i)}2")
        ws.cell(r, 2, f"=IFERROR(AVERAGE({rango(i)})*252,0)").number_format = PCT
        ws.cell(r, 3, f"=IFERROR(_xlfn.STDEV.S({rango(i)})*SQRT(252),0)").number_format = PCT
        ws.cell(r, 4, f'=IFERROR((B{r}-$B$3)/C{r},"")').number_format = "0.000"
        _entrada(ws.cell(r, 5, pesos[i]), PCT)
        ws.cell(r, 6, "=$B$5").number_format = PCT
        ws.cell(r, 7, "=$B$6").number_format = PCT
    etiquetas = nombres if precios is not None else [f"='{datos.title}'!{L(2 + k)}2" for k in range(n)]
    f_cov = f_e + n + 2
    ws.cell(f_cov - 1, 1, "Matriz de covarianzas anual = COVARIANCE.S × 252").font = SECCION
    fc, cc = _matriz(ws, f_cov, 1, etiquetas,
                     lambda i, j: f"=IFERROR(_xlfn.COVARIANCE.S({rango(i)},{rango(j)})*252,0)")
    f_cor = f_cov + n + 3
    ws.cell(f_cor - 1, 1, "Matriz de correlaciones").font = SECCION
    _matriz(ws, f_cor, 1, etiquetas,
            lambda i, j: f'=IFERROR(CORREL({rango(i)},{rango(j)}),"")')
    f_res = f_cor + n + 3
    celdas = _resultados(ws, f_res, n, f_e, "B", "E", fc, cc, "$B$3", "$B$4")
    if res:
        _frontera(ws, f_res + 2 * n + 14, res, f"K{f_e}")
    ws.column_dimensions["A"].width = 42
    for c in range(2, n + 12):
        ws.column_dimensions[L(c)].width = 14
    ws.freeze_panes = "A3"
    return {"hoja": ws.title, "datos": datos.title, "ultima": ultima, "n": n, **celdas}


def _hoja_parametrica(wb: Workbook) -> dict:
    nombres, e, s, rho = ACELERA
    n = len(nombres)
    ws = wb.create_sheet("Paramétrico Acelera")
    _encabezado(ws, "Ejercicio del Taller de Portafolios (Acelera)", 8,
                "Matriz de covarianzas armada con correlaciones. Celdas azul claro: usted escribe. " + AVISO)
    ws["A3"], ws["B3"] = "Tasa libre de riesgo r_f", 0.07
    ws["A4"], ws["B4"] = "Aversión al riesgo A", 4
    ws["A3"].font = ws["A4"].font = NEGRITA
    _entrada(ws["B3"], PCT)
    _entrada(ws["B4"], "0.0")
    f_e = 7
    for k, t in enumerate(["Activo", "E(r) anual", "σ anual", "Sharpe", "Peso w"]):
        _cabecera(ws.cell(f_e - 1, 1 + k, t))
    for i in range(n):
        r = f_e + i
        ws.cell(r, 1, nombres[i])
        for col, v in ((2, e[i]), (3, s[i]), (5, round(1 / n, 6))):
            _entrada(ws.cell(r, col, v), PCT)
        ws.cell(r, 4, f'=IFERROR((B{r}-$B$3)/C{r},"")').number_format = "0.000"
    f_rho = f_e + n + 2
    ws.cell(f_rho - 1, 1, "Correlaciones ρ (cambie el triángulo azul; el otro se copia)").font = SECCION
    fr, cr = _matriz(ws, f_rho, 1, nombres, lambda i, j: 1 if i == j else (rho[i][j] if j > i else f"={L(2 + i)}{f_rho + 1 + j}"))
    for i in range(n):
        for j in range(i + 1, n):
            _entrada(ws.cell(fr + i, cr + j))
    f_cov = f_rho + n + 3
    ws.cell(f_cov - 1, 1, "Covarianzas cov_ij = ρ_ij · σ_i · σ_j").font = SECCION
    fc, cc = _matriz(ws, f_cov, 1, nombres, lambda i, j: f"={L(cr + j)}{fr + i}*$C${f_e + i}*$C${f_e + j}")
    celdas = _resultados(ws, f_cov + n + 3, n, f_e, "B", "E", fc, cc, "$B$3", "$B$4")
    # Dos activos con correlación variable (preguntas 1–3 del ejercicio original, sin las referencias rotas)
    f2 = f_cov + 3 * n + 16
    ws.cell(f2, 1, "Dos activos (A = IPC, B = SP500): σ_p = √(wA²σA² + wB²σB² + 2·wA·wB·ρ·σA·σB)").font = SECCION
    ws.cell(f2 + 1, 1, "ρ(A,B) para esta tabla").font = NEGRITA
    _entrada(ws.cell(f2 + 1, 2, 0.5))
    for k, t in enumerate(["wA", "wB", "σ_p", "E(r_p)", "σ_p con ρ=+1", "σ_p con ρ=0", "σ_p con ρ=−1"]):
        _cabecera(ws.cell(f2 + 2, 1 + k, t))
    for k in range(11):
        r = f2 + 3 + k
        ws.cell(r, 1, k / 10).number_format = PCT
        ws.cell(r, 2, f"=1-A{r}").number_format = PCT
        sa, sb, ea, eb = f"$C${f_e}", f"$C${f_e + 1}", f"$B${f_e}", f"$B${f_e + 1}"
        base = f"A{r}^2*{sa}^2+B{r}^2*{sb}^2+2*A{r}*B{r}*{{rho}}*{sa}*{sb}"
        ws.cell(r, 3, f"=SQRT(MAX({base.format(rho=f'$B${f2 + 1}')},0))").number_format = PCT
        ws.cell(r, 4, f"=A{r}*{ea}+B{r}*{eb}").number_format = PCT
        for col, rv in ((5, "1"), (6, "0"), (7, "-1")):
            ws.cell(r, col, f"=SQRT(MAX({base.format(rho=rv)},0))").number_format = PCT
    g = ScatterChart()
    g.title, g.style, g.height, g.width = "Dos activos: efecto de la correlación", 13, 9, 16
    g.x_axis.title, g.y_axis.title = "σ_p", "E(r_p)"
    ys = Reference(ws, min_col=4, min_row=f2 + 3, max_row=f2 + 13)
    for col, nm, color in ((3, "ρ elegida", MARCA), (5, "ρ=+1", "A3261C"), (6, "ρ=0", SERIE_2), (7, "ρ=−1", POSITIVO)):
        se = Series(ys, Reference(ws, min_col=col, min_row=f2 + 3, max_row=f2 + 13), title=nm)
        _colorear(se, color)
        g.series.append(se)
    ws.add_chart(g, f"I{f2}")
    ws.column_dimensions["A"].width = 46
    for c in range(2, 9):
        ws.column_dimensions[L(c)].width = 14
    ws.freeze_panes = "A3"
    return {"hoja": ws.title, "rho": f"{L(cr + 1)}{fr}:{L(cr + n - 1)}{fr + n - 2}", "f_e": f_e, "n": n, **celdas}


def _pasos_solver(h: dict) -> list[str]:
    pesos, hoja = h["pesos"], h["hoja"]
    return [
        "Active Solver una sola vez: Archivo → Opciones → Complementos → abajo «Administrar: Complementos de Excel» → Ir → marque "
        "«Solver» → Aceptar.",
        f"Vaya a la hoja «{hoja}» y abra Datos → Solver.",
        f"Establecer objetivo: {h['Sharpe']} (Sharpe) con «Máx.»  —  o {h['σ_p anual']} (riesgo σ_p) con «Mín.» si busca el de "
        "menor riesgo.",
        f"Cambiando las celdas de variables: {pesos}.",
        f"Restricciones (botón Agregar): {h['Suma de pesos (debe ser 100 %)']} = 1  ·  {pesos} <= $B$5  ·  {pesos} >= $B$6.",
        "Método de resolución: GRG Nonlinear → Resolver → «Conservar solución de Solver». Los resultados verdes se actualizan solos.",
    ]


def _bytes(wb: Workbook) -> bytes:
    _fuente_terminal(wb)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def plantilla() -> bytes:
    wb = Workbook()
    p = _hoja_parametrica(wb)
    h = _hoja_historica(wb, "Histórico ", [f"Activo {k}" for k in range(1, 6)], None, 253, [0.2] * 5, 0.07, 4, 0.4, 0, None)
    f_e, n = p["f_e"], p["n"]
    _guia(wb, "Plantilla del modelo de Markowitz", [
        ("¿Para qué sirve?", ["Calcula qué mezcla de inversiones da el mejor rendimiento por cada unidad de riesgo (teoría de "
                              "Markowitz). Tiene dos formas de usarla: un ejercicio rápido con supuestos y un modelo con sus "
                              "propios precios."]),
        ("Opción A · Ejercicio rápido (hoja «Paramétrico Acelera»)", [
            "B3: tasa libre de riesgo anual (use la de CETES; p. ej. 7 %). B4: aversión al riesgo A (2 agresivo · 4–6 moderado · "
            "10 conservador).",
            f"B{f_e}:C{f_e + n - 1}: rendimiento esperado y riesgo (σ) anual de cada activo. E{f_e}:E{f_e + n - 1}: los pesos "
            "(deben sumar 100 %).",
            f"{p['rho']}: correlaciones entre activos (triángulo azul; el resto se copia solo).",
            f"Lea el resultado: rendimiento {p['E(r_p) anual']}, riesgo {p['σ_p anual']} y Sharpe {p['Sharpe']}. Más abajo, la "
            "tabla y la gráfica de dos activos muestran cómo la correlación baja el riesgo."]),
        ("Opción B · Con sus precios (hojas «Histórico Datos» e «Histórico Modelo»)", [
            "«Histórico Datos», fila 2 (B2:F2): escriba el nombre de sus 5 activos.",
            "Desde la fila 3: fecha en la columna A y precio de cierre de cada activo en B–F, del más antiguo al más reciente "
            "(ideal: 252 días, un año). Puede usar menos filas: las vacías no cuentan.",
            "Los rendimientos diarios (columnas H–L) y toda la hoja «Histórico Modelo» se calculan solos.",
            f"«Histórico Modelo»: B3 tasa libre de riesgo, B4 aversión A, B5 y B6 peso máximo y mínimo, {h['pesos']} sus pesos.",
            f"Resultado: rendimiento {h['E(r_p) anual']}, riesgo {h['σ_p anual']}, Sharpe {h['Sharpe']}. La suma de pesos "
            f"({h['Suma de pesos (debe ser 100 %)']}) debe dar 100 %."]),
        ("Encontrar la mejor mezcla con Solver (Opción B)", _pasos_solver(h)),
        ("Cómo leer los resultados", [
            "E(r_p): rendimiento anual promedio que habría tenido la mezcla. σ_p: cuánto varía (riesgo). Sharpe: rendimiento extra "
            "sobre la tasa libre de riesgo por unidad de riesgo; más alto es mejor.",
            "y*: qué parte de su dinero pondría en esta mezcla según su aversión A; el resto iría a la tasa libre de riesgo (p. ej. "
            "CETES).",
            "Son datos del pasado: no garantizan resultados futuros. Compare siempre contra repartir en partes iguales."]),
    ])
    return _bytes(wb)


def libro(con, ajustes, portafolios: list[dict]) -> bytes:
    wb = Workbook()
    resumen = wb.create_sheet("Resumen")
    _encabezado(resumen, "Resumen de portafolios", 11, "Cifras enlazadas a cada hoja «P# Modelo»: si cambia los pesos allá, aquí se "
                                                        "actualizan. " + AVISO)
    cab = ["#", "Portafolio", "Perfil", "Método", "Activos", "Ventana", "E(r_p)", "σ_p", "Sharpe", "Capital", "Hoja"]
    for k, t in enumerate(cab, start=1):
        _cabecera(resumen.cell(3, k, t))
    hojas = []
    for k, p in enumerate(portafolios, start=1):
        s = p["especificacion"]
        res = markowitz.calcular(con, ajustes, s)  # se recalcula con los datos actuales para que las hojas cuadren
        precios, _ = markowitz._datos(con, ajustes, [a["id"] for a in res["activos"]], s["ventana"])
        claves = [a["clave"] for a in res["activos"]]
        celdas = _hoja_historica(wb, f"P{k} ", claves, precios, 0, [a["peso"] for a in res["activos"]], s["rf"],
                                 s["aversion"], s["peso_max"], s["peso_min"], res)
        hojas.append((s["nombre"], celdas))
        h = f"'{celdas['hoja']}'"
        for c, v in enumerate([k, s["nombre"], res["objetivo_nombre"], res["metodo_nombre"], ", ".join(claves), s["ventana"],
                               f"={h}!{celdas['E(r_p) anual']}", f"={h}!{celdas['σ_p anual']}", f"={h}!{celdas['Sharpe']}",
                               s["capital"], celdas["hoja"]], start=1):
            cel = resumen.cell(3 + k, c, v)
            cel.number_format = PCT if c in (7, 8) else ("0.000" if c == 9 else (NUM if c == 10 else "General"))
            cel.alignment = Alignment(wrap_text=c == 5, vertical="top")
            cel.border = Border(bottom=LINEA)
    for c, w in zip("ABCDEFGHIJK", (4, 30, 16, 34, 40, 9, 10, 10, 9, 14, 14)):
        resumen.column_dimensions[c].width = w
    resumen.freeze_panes = "A4"
    _guia(wb, f"{len(portafolios)} portafolio(s) del generador", [
        ("Qué contiene", [
            "«Resumen»: los portafolios lado a lado (rendimiento, riesgo y Sharpe enlazados a sus hojas).",
            "Por cada portafolio, «P# Datos» (precios en MXN y rendimientos diarios) y «P# Modelo» (estadísticos, covarianzas, "
            "correlaciones, pesos, resultados, frontera eficiente y gráfica).",
            "Los pesos en azul claro son los que calculó la terminal. Cámbielos para explorar: todo se recalcula."]),
        ("Dónde está cada cosa", [f"«{nombre}» → hoja «{c['hoja']}»: pesos {c['pesos']} · rendimiento {c['E(r_p) anual']} · riesgo "
                                  f"{c['σ_p anual']} · Sharpe {c['Sharpe']} · suma de pesos {c['Suma de pesos (debe ser 100 %)']}."
                                  for nombre, c in hojas]),
        ("Volver a optimizar con Solver", _pasos_solver(hojas[0][1]) + (
            ["En los demás portafolios use las celdas de su hoja (lista de arriba)."] if len(hojas) > 1 else [])),
        ("Cómo leer los resultados", [
            "E(r_p): rendimiento anual promedio en la ventana. σ_p: riesgo. Sharpe: rendimiento extra por unidad de riesgo (más alto "
            "es mejor). La gráfica muestra dónde queda su portafolio frente a la frontera eficiente.",
            "Son datos del pasado: no garantizan resultados futuros."]),
    ])
    return _bytes(wb)
